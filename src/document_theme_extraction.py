from pathlib import Path
import pdfplumber
from docx import Document
from pptx import Presentation
import pandas as pd

# -----------------------------
# DEFINE PATHS
# -----------------------------
BASE_PATH = Path(__file__).resolve().parents[1]
DOC_PATH = BASE_PATH / "data" / "raw" / "Unstructured"
OUTPUT_PATH = BASE_PATH / "data" / "document_intelligence"

OUTPUT_PATH.mkdir(parents=True, exist_ok=True)

# -----------------------------
# DEFINE THEMES + KEYWORDS
# -----------------------------
THEMES = {
    "engagement": ["engagement", "motivation", "satisfaction"],
    "wellbeing": ["burnout", "stress", "wellbeing", "absence"],
    "learning": ["training", "learning", "skills", "development"],
    "inclusion": ["diversity", "inclusion", "equity"],
    "performance": ["performance", "productivity", "efficiency"],
    "governance": ["governance", "compliance", "risk"],
    "finance": ["cost", "fte", "budget", "revenue"]
}

# -----------------------------
# TEXT EXTRACTION FUNCTIONS
# -----------------------------
def extract_pdf_text(file_path):
    text = ""
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text += page.extract_text() or ""
    except:
        pass
    return text.lower()

def extract_docx_text(file_path):
    text = ""
    try:
        doc = Document(file_path)
        for para in doc.paragraphs:
            text += para.text + " "
    except:
        pass
    return text.lower()

def extract_pptx_text(file_path):
    text = ""
    try:
        prs = Presentation(file_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + " "
    except:
        pass
    return text.lower()

# -----------------------------
# PROCESS DOCUMENTS
# -----------------------------
results = []

for file in DOC_PATH.rglob("*"):
    if file.suffix.lower() not in [".pdf", ".docx", ".pptx"]:
        continue

    print(f"Processing {file.name}")

    if file.suffix.lower() == ".pdf":
        text = extract_pdf_text(file)
    elif file.suffix.lower() == ".docx":
        text = extract_docx_text(file)
    elif file.suffix.lower() == ".pptx":
        text = extract_pptx_text(file)

    for theme, keywords in THEMES.items():
        count = sum(text.count(k) for k in keywords)

        if count > 0:
            results.append({
                "document": file.name,
                "theme": theme,
                "keyword_count": count
            })

# -----------------------------
# CREATE DOCUMENT INVENTORY
# -----------------------------
inventory = []

for file in DOC_PATH.rglob("*"):
    if file.suffix.lower() not in [".pdf", ".docx", ".pptx"]:
        continue

    relative_path = file.relative_to(DOC_PATH)

    inventory.append({
        "document": file.name,
        "file_type": file.suffix.lower().replace(".", ""),
        "folder": str(relative_path.parent),
        "path": str(relative_path),
        "size_kb": round(file.stat().st_size / 1024, 2),
    })

inventory_df = pd.DataFrame(inventory)
inventory_df.to_csv(OUTPUT_PATH / "document_inventory.csv", index=False)

# -----------------------------
# SAVE OUTPUT
# -----------------------------
df = pd.DataFrame(results)

df.to_csv(OUTPUT_PATH / "document_theme_summary.csv", index=False)

print("✅ document_theme_summary.csv created!")